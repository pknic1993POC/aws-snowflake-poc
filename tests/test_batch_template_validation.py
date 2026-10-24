import sys
from pathlib import Path

from openpyxl import Workbook

SRC_DIR = Path(__file__).resolve().parents[1] / "ia-converter" / "src"
sys.path.insert(0, str(SRC_DIR))

from converter import validate_batch_template


def test_validate_batch_template_accepts_interface_reference_with_numbered_sheet_name(tmp_path):
    file_path = tmp_path / "batch.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet 1"
    workbook.create_sheet("Interface Reference 2")
    workbook.save(file_path)

    result = validate_batch_template(file_path)

    assert result["valid"] is True
    assert result["errors"] == []


def test_validate_batch_template_requires_interface_reference_sheet(tmp_path):
    file_path = tmp_path / "batch.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet 1"
    workbook.create_sheet("Other Sheet")
    workbook.save(file_path)

    result = validate_batch_template(file_path)

    assert result["valid"] is False
    assert any("Interface Reference" in error for error in result["errors"])
